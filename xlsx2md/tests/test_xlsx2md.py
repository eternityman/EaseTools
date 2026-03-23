# -*- coding: utf-8 -*-
"""
test_xlsx2md.py - xlsx2md 单元测试

使用 pytest 运行：
    pip install pytest openpyxl Pillow
    pytest tests/
"""

import io
import os
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.drawing.spreadsheet_drawing import (
    OneCellAnchor,
    TwoCellAnchor,
)
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU

# 将父目录添加到 sys.path，使测试可以直接导入 xlsx2md
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from xlsx2md import (
    cell_value_to_str,
    get_anchor_cell,
    sheet_to_markdown,
    xlsx_to_markdown,
    xlsx_to_markdown_string,
)


class TestGetAnchorCell(unittest.TestCase):
    """测试锚点位置解析函数"""

    def test_two_cell_anchor(self):
        """TwoCellAnchor 应返回 _from 的 (row, col)"""
        anchor = TwoCellAnchor()
        anchor._from.row = 2
        anchor._from.col = 3
        self.assertEqual(get_anchor_cell(anchor), (2, 3))

    def test_one_cell_anchor(self):
        """OneCellAnchor 应返回 _from 的 (row, col)"""
        anchor = OneCellAnchor()
        anchor._from.row = 0
        anchor._from.col = 1
        self.assertEqual(get_anchor_cell(anchor), (0, 1))

    def test_unknown_anchor(self):
        """未知锚点类型应返回 None"""
        self.assertIsNone(get_anchor_cell(object()))

    def test_none_anchor(self):
        """None 锚点应返回 None"""
        self.assertIsNone(get_anchor_cell(None))


class TestCellValueToStr(unittest.TestCase):
    """测试单元格值转字符串函数"""

    def _make_cell(self, value):
        """创建一个带有指定值的 openpyxl Cell"""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = value
        return ws["A1"]

    def test_none_returns_empty(self):
        cell = self._make_cell(None)
        self.assertEqual(cell_value_to_str(cell), "")

    def test_string_value(self):
        cell = self._make_cell("hello")
        self.assertEqual(cell_value_to_str(cell), "hello")

    def test_integer_value(self):
        cell = self._make_cell(42)
        self.assertEqual(cell_value_to_str(cell), "42")

    def test_newline_replaced_with_br(self):
        cell = self._make_cell("line1\nline2")
        self.assertEqual(cell_value_to_str(cell), "line1<br>line2")

    def test_cr_lf_replaced_with_br(self):
        cell = self._make_cell("line1\r\nline2")
        self.assertEqual(cell_value_to_str(cell), "line1<br>line2")

    def test_pipe_escaped(self):
        cell = self._make_cell("a|b")
        self.assertEqual(cell_value_to_str(cell), r"a\|b")


class TestSheetToMarkdown(unittest.TestCase):
    """测试工作表转 Markdown 函数"""

    def _make_sheet(self, data):
        """根据二维列表创建 openpyxl 工作表"""
        wb = openpyxl.Workbook()
        ws = wb.active
        for r_idx, row in enumerate(data, start=1):
            for c_idx, val in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=val)
        return ws

    def test_basic_table(self):
        ws = self._make_sheet([
            ["姓名", "年龄"],
            ["张三", 28],
            ["李四", 32],
        ])
        lines = sheet_to_markdown(ws, {})
        self.assertEqual(lines[0], "| 姓名 | 年龄 |")
        self.assertIn("---", lines[1])
        self.assertEqual(lines[2], "| 张三 | 28 |")
        self.assertEqual(lines[3], "| 李四 | 32 |")

    def test_empty_sheet(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        lines = sheet_to_markdown(ws, {})
        # 空 Sheet 应返回提示信息而非崩溃
        self.assertTrue(len(lines) > 0)

    def test_image_reference_inserted(self):
        ws = self._make_sheet([
            ["姓名", "头像"],
            ["张三", None],
        ])
        # 行 1（0-based），列 1（0-based）对应第 2 行第 2 列
        img_map = {(1, 1): ["images/Sheet1_img_0.png"]}
        lines = sheet_to_markdown(ws, img_map)
        self.assertIn("![image](images/Sheet1_img_0.png)", lines[2])

    def test_image_and_text_combined(self):
        ws = self._make_sheet([
            ["名称", "内容"],
            ["测试", "备注文字"],
        ])
        # (1, 1) 对应 "备注文字" 单元格
        img_map = {(1, 1): ["images/img.png"]}
        lines = sheet_to_markdown(ws, img_map)
        # 文字和图片引用都应出现
        self.assertIn("备注文字", lines[2])
        self.assertIn("![image](images/img.png)", lines[2])

    def test_no_images_flag(self):
        ws = self._make_sheet([
            ["A", "B"],
            ["1", "2"],
        ])
        img_map = {(1, 0): ["images/img.png"]}
        lines = sheet_to_markdown(ws, img_map, include_images=False)
        # 禁用图片时，Markdown 中不应有图片引用
        full_text = "\n".join(lines)
        self.assertNotIn("![image]", full_text)


class TestXlsxToMarkdown(unittest.TestCase):
    """测试完整的 XLSX → Markdown 转换函数（集成测试）"""

    def _create_xlsx(self, data_by_sheet):
        """创建含多 Sheet 的 XLSX 文件，返回临时文件路径"""
        wb = openpyxl.Workbook()
        # 移除默认 Sheet
        wb.remove(wb.active)
        for sheet_name, rows in data_by_sheet.items():
            ws = wb.create_sheet(title=sheet_name)
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        tmp.close()
        return tmp.name

    def test_basic_conversion(self):
        """基本表格转换，验证 Markdown 输出包含正确的内容"""
        xlsx_path = self._create_xlsx({
            "Sheet1": [
                ["姓名", "职位"],
                ["张三", "经理"],
                ["李四", "工程师"],
            ]
        })
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, include_images=False)
                self.assertTrue(os.path.isfile(md_path))
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                self.assertIn("## Sheet1", content)
                self.assertIn("姓名", content)
                self.assertIn("张三", content)
                self.assertIn("工程师", content)
        finally:
            os.unlink(xlsx_path)

    def test_multi_sheet(self):
        """多 Sheet 转换，每个 Sheet 都应出现在输出中"""
        xlsx_path = self._create_xlsx({
            "Alpha": [["A", "B"], [1, 2]],
            "Beta":  [["X", "Y"], [3, 4]],
        })
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, include_images=False)
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                self.assertIn("## Alpha", content)
                self.assertIn("## Beta", content)
        finally:
            os.unlink(xlsx_path)

    def test_sheet_filter(self):
        """--sheet 参数只转换指定 Sheet"""
        xlsx_path = self._create_xlsx({
            "Alpha": [["A"], [1]],
            "Beta":  [["B"], [2]],
        })
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(
                    xlsx_path,
                    output_path=md_path,
                    include_images=False,
                    sheet_filter="Alpha",
                )
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                self.assertIn("## Alpha", content)
                self.assertNotIn("## Beta", content)
        finally:
            os.unlink(xlsx_path)

    def test_no_images_flag(self):
        """--no-images 标志不应产生图片引用"""
        xlsx_path = self._create_xlsx({
            "Sheet1": [["Name"], ["Test"]],
        })
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, include_images=False)
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                self.assertNotIn("![image]", content)
        finally:
            os.unlink(xlsx_path)

    def test_default_output_path(self):
        """默认输出路径应与输入同名，扩展名改为 .md"""
        xlsx_path = self._create_xlsx({"S": [["A"], [1]]})
        expected_md = os.path.splitext(xlsx_path)[0] + ".md"
        try:
            xlsx_to_markdown(xlsx_path, include_images=False)
            self.assertTrue(os.path.isfile(expected_md))
        finally:
            os.unlink(xlsx_path)
            if os.path.isfile(expected_md):
                os.unlink(expected_md)

    def test_missing_input_file(self):
        """输入文件不存在时应调用 sys.exit(1)"""
        with self.assertRaises(SystemExit) as ctx:
            xlsx_to_markdown("/nonexistent/path/file.xlsx")
        self.assertEqual(ctx.exception.code, 1)


class TestXlsxToMarkdownString(unittest.TestCase):
    """测试 xlsx_to_markdown_string 函数（返回字符串而不写入文件）"""

    def _create_xlsx(self, data_by_sheet):
        """创建含多 Sheet 的 XLSX 文件，返回临时文件路径"""
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet_name, rows in data_by_sheet.items():
            ws = wb.create_sheet(title=sheet_name)
            for r_idx, row in enumerate(rows, start=1):
                for c_idx, val in enumerate(row, start=1):
                    ws.cell(row=r_idx, column=c_idx, value=val)
        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        tmp.close()
        return tmp.name

    def test_returns_string(self):
        """返回值应为字符串类型"""
        xlsx_path = self._create_xlsx({"Sheet1": [["A"], [1]]})
        try:
            result = xlsx_to_markdown_string(xlsx_path, include_images=False)
            self.assertIsInstance(result, str)
        finally:
            os.unlink(xlsx_path)

    def test_content_matches_file_output(self):
        """返回的字符串内容应与写入文件的内容一致"""
        xlsx_path = self._create_xlsx({
            "Data": [["名称", "数量"], ["苹果", 10], ["香蕉", 20]],
        })
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, include_images=False)
                with open(md_path, encoding="utf-8") as f:
                    file_content = f.read()
            string_content = xlsx_to_markdown_string(xlsx_path, include_images=False)
            self.assertEqual(string_content, file_content)
        finally:
            os.unlink(xlsx_path)

    def test_missing_file_raises(self):
        """输入文件不存在时应抛出 FileNotFoundError"""
        with self.assertRaises(FileNotFoundError):
            xlsx_to_markdown_string("/nonexistent/path/file.xlsx")

    def test_sheet_filter(self):
        """sheet_filter 参数只返回指定 Sheet 内容"""
        xlsx_path = self._create_xlsx({
            "Alpha": [["A"], [1]],
            "Beta": [["B"], [2]],
        })
        try:
            result = xlsx_to_markdown_string(xlsx_path, include_images=False, sheet_filter="Alpha")
            self.assertIn("## Alpha", result)
            self.assertNotIn("## Beta", result)
        finally:
            os.unlink(xlsx_path)

    def test_empty_result_for_unknown_sheet(self):
        """sheet_filter 指定不存在的 Sheet 时返回空字符串"""
        xlsx_path = self._create_xlsx({"Sheet1": [["A"], [1]]})
        try:
            result = xlsx_to_markdown_string(xlsx_path, include_images=False, sheet_filter="NonExistent")
            self.assertEqual(result, "")
        finally:
            os.unlink(xlsx_path)

    def test_images_embedded_in_string(self):
        """当 XLSX 含嵌入图片时，返回的字符串中应含图片引用"""
        try:
            from PIL import Image as PILImage
            from openpyxl.drawing.image import Image as XLImage
        except ImportError:
            self.skipTest("Pillow 或 openpyxl 未安装")

        buf = io.BytesIO()
        PILImage.new("RGB", (10, 10), color=(255, 0, 0)).save(buf, format="PNG")
        buf.seek(0)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "Name"
        ws["B1"] = "Photo"
        ws["A2"] = "Alice"
        xl_img = XLImage(buf)
        xl_img.anchor = "B2"
        ws.add_image(xl_img)

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        tmp.close()
        try:
            result = xlsx_to_markdown_string(tmp.name)
            self.assertIn("![image]", result)
        finally:
            os.unlink(tmp.name)


class TestImagePlacement(unittest.TestCase):
    """
    真实端到端图片测试：
    将实际 PNG 图片嵌入 XLSX 特定单元格，
    运行完整转换，验证 Markdown 输出中图片引用出现在正确的行和列。
    """

    @staticmethod
    def _make_png(color=(255, 0, 0)):
        """生成一个 10×10 纯色 PNG 的 BytesIO"""
        try:
            from PIL import Image as PILImage
        except ImportError:
            return None
        buf = io.BytesIO()
        PILImage.new("RGB", (10, 10), color=color).save(buf, format="PNG")
        buf.seek(0)
        return buf

    def _create_xlsx_with_images(self):
        """
        创建包含嵌入图片的 XLSX：
          A1=姓名, B1=头像, C1=备注
          A2=张三,          C2=经理   （图片锚定在 B2）
          A3=李四,          C3=工程师  （图片锚定在 B3）
        返回临时 XLSX 路径。
        """
        try:
            from openpyxl.drawing.image import Image as XLImage
        except ImportError:
            self.skipTest("openpyxl 未安装")

        png1 = self._make_png((255, 0, 0))
        png2 = self._make_png((0, 0, 255))
        if png1 is None:
            self.skipTest("Pillow 未安装，跳过图片测试")

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "姓名"
        ws["B1"] = "头像"
        ws["C1"] = "备注"
        ws["A2"] = "张三"
        ws["C2"] = "经理"
        ws["A3"] = "李四"
        ws["C3"] = "工程师"

        img1 = XLImage(png1)
        img1.anchor = "B2"  # 0-based: row=1, col=1
        ws.add_image(img1)

        img2 = XLImage(png2)
        img2.anchor = "B3"  # 0-based: row=2, col=1
        ws.add_image(img2)

        tmp = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        wb.save(tmp.name)
        tmp.close()
        return tmp.name

    def _parse_table_rows(self, md_content):
        """从 Markdown 内容中解析出表格行（跳过分隔行），返回每行各列列表"""
        rows = []
        for line in md_content.splitlines():
            if not line.startswith("|"):
                continue
            # 跳过分隔行 (| --- | --- |)
            stripped = line.strip("|").strip()
            if all(set(c) <= set("- ") for c in stripped.split("|")):
                continue
            cols = [c.strip() for c in line.strip("|").split("|")]
            rows.append(cols)
        return rows

    def test_images_saved_to_disk(self):
        """图片文件应被实际保存到 images 目录"""
        xlsx_path = self._create_xlsx_with_images()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, image_dir="images")
                img_dir = os.path.join(tmpdir, "images")
                self.assertTrue(os.path.isdir(img_dir), "images 目录应存在")
                saved = sorted(os.listdir(img_dir))
                self.assertEqual(len(saved), 2, f"应保存 2 张图片，实际: {saved}")
                for fname in saved:
                    fpath = os.path.join(img_dir, fname)
                    self.assertGreater(os.path.getsize(fpath), 0, f"{fname} 不应为空文件")
        finally:
            os.unlink(xlsx_path)

    def test_image_in_correct_row_and_column(self):
        """
        图片引用应出现在 Markdown 表格中正确的行和列：
        - 张三行（数据第 1 行）的 B 列（第 2 列）包含图片引用
        - 李四行（数据第 2 行）的 B 列（第 2 列）包含图片引用
        - 备注列（第 3 列）不应含图片引用
        """
        xlsx_path = self._create_xlsx_with_images()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, image_dir="images")
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                rows = self._parse_table_rows(content)

                # rows[0] = 表头行
                self.assertEqual(rows[0][0], "姓名")
                self.assertEqual(rows[0][1], "头像")
                self.assertEqual(rows[0][2], "备注")

                # rows[1] = 张三行
                self.assertEqual(rows[1][0], "张三", "第1列应为张三")
                self.assertIn("![image]", rows[1][1], "张三行 B 列应含图片引用")
                self.assertNotIn("![image]", rows[1][2], "张三行 C 列不应含图片引用")
                self.assertIn("经理", rows[1][2], "张三行 C 列应为备注文字")

                # rows[2] = 李四行
                self.assertEqual(rows[2][0], "李四", "第1列应为李四")
                self.assertIn("![image]", rows[2][1], "李四行 B 列应含图片引用")
                self.assertNotIn("![image]", rows[2][2], "李四行 C 列不应含图片引用")
                self.assertIn("工程师", rows[2][2], "李四行 C 列应为备注文字")

        finally:
            os.unlink(xlsx_path)

    def test_two_images_reference_different_files(self):
        """两张图片应引用不同的文件名"""
        xlsx_path = self._create_xlsx_with_images()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, image_dir="images")
                with open(md_path, encoding="utf-8") as f:
                    content = f.read()
                rows = self._parse_table_rows(content)

                import re
                refs_row1 = re.findall(r'!\[image\]\(([^)]+)\)', rows[1][1])
                refs_row2 = re.findall(r'!\[image\]\(([^)]+)\)', rows[2][1])
                self.assertEqual(len(refs_row1), 1)
                self.assertEqual(len(refs_row2), 1)
                self.assertNotEqual(
                    refs_row1[0], refs_row2[0],
                    "两行图片应引用不同的文件"
                )
        finally:
            os.unlink(xlsx_path)

    def test_image_file_names_use_sheet_name(self):
        """图片文件名应包含 Sheet 名称作为前缀"""
        xlsx_path = self._create_xlsx_with_images()
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                md_path = os.path.join(tmpdir, "output.md")
                xlsx_to_markdown(xlsx_path, output_path=md_path, image_dir="images")
                img_dir = os.path.join(tmpdir, "images")
                for fname in os.listdir(img_dir):
                    self.assertTrue(
                        fname.startswith("Sheet1"),
                        f"图片文件名应以 Sheet1 开头，实际: {fname}"
                    )
        finally:
            os.unlink(xlsx_path)


if __name__ == "__main__":
    unittest.main()
